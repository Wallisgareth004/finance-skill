"""Write-back tests — run against a COPY of the prototype, never a live file.

The contract under test:
    backup → write to temp → verify → atomic replace
    and on ANY failure, the original is restored untouched.
"""

import os
import shutil
import tempfile
import unittest
import zipfile
from datetime import date
from decimal import Decimal

from helpers import pipeline_dir, prototype_workbook

pipeline_dir()
import writer  # noqa: E402

PROTOTYPE = prototype_workbook()  # see helpers.prototype_workbook (2026-08-25)


class _Row:
    def __init__(self, d, desc, cat, amt, bal, slug):
        self.date = d
        self.description = desc
        self.category = cat
        self.amount = Decimal(amt)
        self.balance = Decimal(bal) if bal is not None else None
        self.source_slug = slug


def _rows():
    return [
        _Row(
            date(2026, 8, 20),
            "POS Purchase Test Merchant",
            "Groceries",
            "-123.45",
            "1357.67",
            "export-2026-08-20_to_08-20",
        ),
        _Row(
            date(2026, 8, 20),
            'Weird & "quoted" <merchant>',
            "Uncategorised",
            "-1.00",
            "1356.67",
            "export-2026-08-20_to_08-20",
        ),
    ]


@unittest.skipUnless(PROTOTYPE and os.path.exists(PROTOTYPE), "prototype workbook not present")
class TestAppend(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.book = os.path.join(self.tmp, "model.xlsx")
        shutil.copy2(PROTOTYPE, self.book)
        # Derived, NOT hardcoded. Row numbers baked into assertions rot the moment
        # anything appends to the fixture — which is exactly what this module does.
        import openpyxl
        wb = openpyxl.load_workbook(self.book, read_only=True)
        self.base_row = wb["TRANSACTIONS"].max_row
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _charts(self, p):
        with zipfile.ZipFile(p) as z:
            return sum(1 for n in z.namelist() if "charts/chart" in n)

    def test_rows_appended_and_charts_survive(self):
        before_charts = self._charts(self.book)
        res = writer.append_transactions(self.book, _rows())
        self.assertEqual(res.rows_added, 2)
        self.assertEqual(res.first_row, self.base_row + 1)
        self.assertEqual(res.last_row, self.base_row + 2)
        self.assertEqual(self._charts(self.book), before_charts)
        self.assertEqual(before_charts, 4)

    def test_backup_created_and_matches_pre_state(self):
        import hashlib

        pre = hashlib.sha256(open(self.book, "rb").read()).hexdigest()
        res = writer.append_transactions(self.book, _rows())
        post_backup = hashlib.sha256(open(res.backup_path, "rb").read()).hexdigest()
        self.assertEqual(pre, post_backup, "backup must capture the PRE-write state")

    def test_values_readable_after_write(self):
        import openpyxl

        writer.append_transactions(self.book, _rows())
        wb = openpyxl.load_workbook(self.book, read_only=True)
        ws = wb["TRANSACTIONS"]
        self.assertEqual(ws.max_row, self.base_row + 2)
        last = list(ws.iter_rows(min_row=self.base_row + 1, max_row=self.base_row + 2, values_only=True))
        self.assertEqual(last[0][1], "POS Purchase Test Merchant")
        self.assertEqual(last[0][2], "Groceries")
        self.assertEqual(float(last[0][3]), -123.45)
        wb.close()

    def test_xml_special_characters_survive(self):
        import openpyxl

        writer.append_transactions(self.book, _rows())
        wb = openpyxl.load_workbook(self.book, read_only=True)
        vals = [
            r[1] for r in wb["TRANSACTIONS"].iter_rows(min_row=self.base_row + 2, values_only=True)
        ]
        self.assertEqual(vals[0], 'Weird & "quoted" <merchant>')
        wb.close()

    def test_dimension_and_autofilter_updated(self):
        writer.append_transactions(self.book, _rows())
        with zipfile.ZipFile(self.book) as z:
            xml = z.read("xl/worksheets/sheet6.xml").decode("utf-8")
        self.assertIn(f'<dimension ref="A1:F{self.base_row + 2}"/>', xml)
        self.assertIn(f'ref="A3:F{self.base_row + 2}"', xml, "stale autofilter must be corrected")

    def test_calc_chain_dropped(self):
        writer.append_transactions(self.book, _rows())
        with zipfile.ZipFile(self.book) as z:
            names = z.namelist()
            ct = z.read("[Content_Types].xml").decode("utf-8")
        self.assertFalse(any("calcChain" in n for n in names))
        self.assertNotIn("calcChain", ct)

    def test_repeated_writes_are_additive_not_corrupting(self):
        import openpyxl

        writer.append_transactions(self.book, _rows())
        writer.append_transactions(self.book, _rows())
        wb = openpyxl.load_workbook(self.book, read_only=True)
        self.assertEqual(wb["TRANSACTIONS"].max_row, self.base_row + 4)
        wb.close()
        self.assertEqual(self._charts(self.book), 4)

    def test_failure_restores_the_original(self):
        import hashlib

        pre = hashlib.sha256(open(self.book, "rb").read()).hexdigest()
        bad = _rows()[0]
        bad.date = "not-a-date"  # excel_serial will raise
        with self.assertRaises(Exception):
            writer.append_transactions(self.book, [bad])
        post = hashlib.sha256(open(self.book, "rb").read()).hexdigest()
        self.assertEqual(
            pre, post, "original must be byte-identical after a failed write"
        )
        root, ext = os.path.splitext(self.book)
        self.assertFalse(os.path.exists(f"{root}.tmp-write{ext}"))

    def test_empty_write_refused(self):
        with self.assertRaises(writer.WriteFailed):
            writer.append_transactions(self.book, [])


class TestBackupRotation(unittest.TestCase):
    def test_rotation_keeps_only_n(self):
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, "f.xlsx")
            with open(p, "wb") as fh:
                fh.write(b"x")
            for _ in range(5):
                writer.backup(p, keep=3)
                import time

                time.sleep(1.01)  # timestamps are second-resolution
            found = [f for f in os.listdir(tmp) if ".backup-" in f]
            self.assertEqual(len(found), 3, found)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
