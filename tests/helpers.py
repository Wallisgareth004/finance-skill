"""Synthetic fixtures mirroring the REAL FNB export formats, byte-for-byte.

Every quirk here was measured from real FNB exports on 2026-08-20. Every VALUE
is synthetic — names, account numbers, dates, amounts and merchants are
invented: the 6-line preamble, the leading space before each field, the
`YYYY/MM/DD` dates, the newest-first row order, and the card's literal `0`
balance column are the only things copied, because they are FNB's format.

The tests are plain unittest, so either runner works from the repository root:
    python3 -m pytest tests
    python3 -m unittest discover -s tests -t .
"""

import os
import sys

PIPELINE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "scripts",
    "finance_pipeline",
)

SKILL_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def pipeline_dir():
    if PIPELINE not in sys.path:
        sys.path.insert(0, PIPELINE)
    return PIPELINE


def live_model():
    """The real workbook, from FINANCE_LIVE_MODEL. None when unset or absent.

    Read-only by contract: every caller copies it before writing.
    """
    path = os.environ.get("FINANCE_LIVE_MODEL")
    return path if path and os.path.exists(path) else None


def real_data(*parts):
    """A path under FINANCE_REAL_DATA_DIR, or None when unset or absent.

    Real-data tests skip on None. No machine layout is written in this file.
    """
    root = os.environ.get("FINANCE_REAL_DATA_DIR")
    if not root:
        return None
    path = os.path.join(root, *parts)
    return path if os.path.exists(path) else None


def prototype_workbook():
    """A real-shaped .xlsx the writer and run tests can copy and mutate.

    Added 2026-08-25 after a validation sweep found 18 writer tests SILENTLY
    SKIPPED: they hardcoded a workbook path that no longer existed, and
    `@unittest.skipUnless` turned every Excel-writer test into a skip. The suite
    reported "228 passed" while the module that writes the workbook had zero
    coverage. A green suite hiding a hole in the highest-risk component is
    exactly what this skill exists to prevent.

    2026-09-24: the live workbook comes from FINANCE_LIVE_MODEL; without it a
    synthetic workbook is built, so the writer tests always RUN. Every caller
    copies the result before writing to it.
    """
    return live_model() or synthetic_workbook()


_SYNTHETIC = None


def synthetic_workbook():
    global _SYNTHETIC
    if _SYNTHETIC is None or not os.path.exists(_SYNTHETIC):
        import tempfile
        from datetime import date

        folder = tempfile.mkdtemp(prefix="finance-synthetic-")
        _SYNTHETIC = make_workbook(
            os.path.join(folder, "synthetic-model.xlsx"),
            rows=[
                (date(2026, 8, 1), "SEED SHOP", "Groceries", "-10.00", "1628.80", "seed"),
                (date(2026, 8, 2), "SEED FEE", "Bank fees & interest", "-3.00", "1625.80", "seed"),
            ],
        )
    return _SYNTHETIC


# NOTE the leading space before Amount/Balance/Description in the header, and
# before every data field. FNB emits them; csv.DictReader would keep them in the
# key names. Everything must be stripped.
CHEQUE_PREAMBLE = (
    "ACCOUNT TRANSACTION HISTORY\n"
    "\n"
    "Name:, Test, User\n"
    "Account:, 10000000001, [Easy Account]\n"
    "Balance:, 2215.40, 615.25\n"
    "\n"
    "Date, Amount, Balance, Description\n"
)

CARD_PREAMBLE = (
    "ACCOUNT TRANSACTION HISTORY\n"
    "\n"
    "Name:, Test, User\n"
    "Account:, 400000******0001, [FNB Aspire Credit Account]\n"
    "Balance:, -1834.52, 915\n"
    "\n"
    "Date, Amount, Balance, Description\n"
)


def make_cheque_csv(directory, rows=None, name="10000000001.csv"):
    """rows: (date, amount, balance, description) as strings, NEWEST FIRST.

    Default rows form a valid balance chain:
        balance[i] == balance[i+1] + amount[i]
    """
    if rows is None:
        rows = [
            (
                "2026/03/12",
                "-54.21",
                "2215.40",
                "210.00 HARBOUR PHAR 400000*0001  09 MAR",
            ),
            (
                "2026/03/12",
                "-61.90",
                "2269.61",
                "WOOLWORTHS          400000*0001  11 MAR",
            ),
            ("2026/03/11", "-3.00", "2331.51", "#INT PYMT FEE-47.13 SPTFY"),
        ]
    path = os.path.join(directory, name)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(CHEQUE_PREAMBLE)
        for d, a, b, desc in rows:
            fh.write(f"{d}, {a}, {b}, {desc}\n")
    return path


def make_card_csv(directory, rows=None, name="400000______0001.csv"):
    """The card's Balance column is the literal string `0` on every row —
    a placeholder, not data. Descriptions are space-padded to exactly 40 chars.
    """
    if rows is None:
        rows = [
            ("2026/03/10", "-138.00", "0", "iK *Harbour Grill       DURBAN          "),
            ("2026/03/09", "-583.20", "0", "SEASIDE TRADING,LDA     Port Harbour    "),
        ]
    path = os.path.join(directory, name)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(CARD_PREAMBLE)
        for d, a, b, desc in rows:
            fh.write(f"{d}, {a}, {b}, {desc}\n")
    return path


# ── synthetic workbook (2026-09-24) ─────────────────────────────────────────
# Same sheet order as the real model — TRANSACTIONS must be the SIXTH sheet,
# because writer.TXNS_SHEET names xl/worksheets/sheet6.xml — with four charts on
# DASHBOARD. Built fresh by openpyxl, which is fine for a NEW synthetic file;
# the no-save rule protects an existing workbook's charts, and this one's charts
# are written by the same library that created them.
#
# Its style table is whatever openpyxl produces, so style index 21 is NOT a date
# format here. That is deliberate: it is the shape of every workbook except the
# live one, and it is how D7 (hardcoded date style) was found.
SHEETS = ("DASHBOARD", "BUDGET", "BY MONTH", "BALANCE", "SUBSCRIPTIONS",
          "TRANSACTIONS", "NOTES & GAPS")


def make_workbook(path, rows=(), months=(("2026-01", 100.0), ("2026-02", 120.0),
                                          ("2026-03", 90.0), ("2026-04", 110.0)),
                  card_rows=None):
    """rows: (date, description, category, amount, balance, source).
    card_rows: (date, description, category, amount) for a CREDIT CARD sheet in
    the real layout (title block, header on row 8, data from row 9, a totals
    row with no date). None = no such sheet, like most workbooks."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in SHEETS}

    by_month = sheets["BY MONTH"]
    by_month.append(["Month", "Spend"])
    for month, value in months:
        by_month.append([month, value])
    sheets["BUDGET"].append(["Line", "Monthly"])
    sheets["BUDGET"].append(["Groceries", 4000])
    sheets["BALANCE"].append(["Date", "Balance"])
    sheets["SUBSCRIPTIONS"].append(["Name", "Monthly", "Decision"])
    sheets["SUBSCRIPTIONS"].append(["Streaming Co", 199, "KEEP"])
    sheets["NOTES & GAPS"].append(["COVERAGE", "synthetic"])
    dash = sheets["DASHBOARD"]
    dash["A1"] = "Synthetic model"
    dash["B2"] = "=SUM('BY MONTH'!B2:B5)"

    data = Reference(by_month, min_col=2, min_row=1, max_row=len(months) + 1)
    cats = Reference(by_month, min_col=1, min_row=2, max_row=len(months) + 1)
    for i, kind in enumerate((BarChart, LineChart, BarChart, LineChart)):
        chart = kind()
        chart.title = f"Chart {i + 1}"
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        dash.add_chart(chart, f"{'DH'[i % 2]}{4 + 16 * (i // 2)}")

    txns = sheets["TRANSACTIONS"]
    txns["A1"] = "TRANSACTIONS"
    for col, name in enumerate(("Date", "Description", "Category", "Amount", "Balance", "Source"), 1):
        txns.cell(row=3, column=col, value=name)
    for r, (d, desc, cat, amt, bal, src) in enumerate(rows, 4):
        txns.cell(row=r, column=1, value=d).number_format = "yyyy-mm-dd"
        txns.cell(row=r, column=2, value=desc)
        txns.cell(row=r, column=3, value=cat)
        txns.cell(row=r, column=4, value=float(amt)).number_format = "#,##0.00"
        if bal is not None:
            txns.cell(row=r, column=5, value=float(bal)).number_format = "#,##0.00"
        txns.cell(row=r, column=6, value=src)
    # The real sheet carries an autoFilter over the header row that the writer
    # must extend on every append (it was found stale once).
    txns.auto_filter.ref = f"A3:F{3 + len(rows)}"
    if card_rows is not None:
        card = wb.create_sheet("CREDIT CARD", index=wb.sheetnames.index("TRANSACTIONS") + 1)
        card["A1"] = "  SYNTHETIC CREDIT CARD"
        card["A3"], card["B3"] = "Balance (owing)", -1234.56
        card["A6"], card["B6"] = "Export window", "synthetic"
        for col, name in enumerate(("Date", "Description", "Category", "Amount"), 1):
            card.cell(row=8, column=col, value=name)
        r = 9
        for r, (d, desc, cat, amt) in enumerate(card_rows, 9):
            card.cell(row=r, column=1, value=d).number_format = "yyyy-mm-dd"
            card.cell(row=r, column=2, value=desc)
            card.cell(row=r, column=3, value=cat)
            card.cell(row=r, column=4, value=float(amt))
        card.cell(row=r + 2, column=2, value="Total")
        card.cell(row=r + 2, column=4, value=float(sum(float(c[3]) for c in card_rows)))
    wb.save(path)
    return path


# ── synthetic "Transaction history" PDF (D11, 2026-09-27) ──────────────────
# FNB's online-banking "Transaction history" download lists card
# authorisations that have not posted yet. Its text, as pypdf's layout mode
# reads it: a header, then per authorisation one line
#     DD Mon YYYY HH:MM:SS    <card mask>    <description>
# and the amount alone on the line below. A page break can carry a row's
# amount to the top of the next page; the real file does, so this does too.
# Written by hand (no PDF library): every value is synthetic.
def _pdf_bytes(pages):
    objs = []

    def add(body):
        objs.append(body)
        return len(objs)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    contents = []
    for texts in pages:
        stream = "".join(
            f"BT /F1 9 Tf {x} {y} Td ({t}) Tj ET\n" for x, y, t in texts
        ).encode("latin-1")
        contents.append(add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream"))
    pages_id = len(objs) + len(pages) + 1
    kids = [
        add(
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] /Contents {c} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        for c in contents
    ]
    add(f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode())
    catalog = add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode())
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root {catalog} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_pending_pdf(path, rows, account="10000000001", card="400000******0001",
                     rows_per_page=20, carry_amount=True, title="Transaction history",
                     omit_amounts=(), stray=None):
    """rows: ("DD Mon YYYY HH:MM:SS", description, amount text), newest last.
    `omit_amounts` (row indexes) and `stray` (a line of text after the rows)
    build the malformed files the parser must refuse."""
    header = [
        (40, 760, title), (40, 745, "Name :"), (250, 745, "Easy"), (40, 735, "Account"),
        (40, 720, "Account Number :"), (250, 720, account), (40, 705, "Type :"),
        (250, 705, "Easy"), (40, 695, "Account"), (40, 680, "Date and Time"),
        (180, 680, "Card Number"), (320, 680, "Description"), (40, 670, "Amount"),
    ]
    chunks = [rows[i:i + rows_per_page] for i in range(0, len(rows), rows_per_page)] or [[]]
    pages, carried = [], None
    for n, chunk in enumerate(chunks):
        texts = list(header) if n == 0 else []
        y = 640 if n == 0 else 760
        if carried is not None:
            texts.append((500, y, carried))
            y -= 20
            carried = None
        for i, (when, desc, amount) in enumerate(chunk):
            texts += [(40, y, when), (180, y, card), (320, y, desc)]
            last_on_page = i == len(chunk) - 1 and n < len(chunks) - 1
            if n * rows_per_page + i in omit_amounts:
                pass
            elif last_on_page and carry_amount:
                carried = amount
            else:
                texts.append((500, y - 12, amount))
            y -= 30
        if stray and n == len(chunks) - 1:
            texts.append((40, y, stray))
        pages.append(texts)
    with open(path, "wb") as fh:
        fh.write(_pdf_bytes(pages))
    return path
