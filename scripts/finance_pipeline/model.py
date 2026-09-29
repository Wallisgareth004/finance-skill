#!/usr/bin/env python3
"""model.py — read the workbook and answer one question: what is genuinely new?

IN   the workbook, read-only
OUT  TRANSACTIONS as a Counter of dedupe keys; what is genuinely new; the
     CREDIT CARD rows

THIS MODULE NEVER WRITES. Saving through openpyxl would delete the workbook's
charts — it does not round-trip xl/charts/*. Reads are read_only=True,
data_only=True. All writes go through writer.append_transactions.

(The literal call spelling is avoided even in this docstring: test_model.py
greps the source for it, and a mention in prose is indistinguishable from a
call to a regex. Keeping the guard dumb keeps it trustworthy.)

═══ WHY A COUNTER, NOT A SET ═══
On the author's real ledger a few percent of rows share their (date,
description, amount) triple with another row. Every such group was checked,
and every one is GENUINE money, not a prior double-import:
  · each group's rows have all-distinct running balances
  · they sit on consecutive links of the balance chain, each row exactly
    |amount| below the last, one chain link per row
  · no group spans two exports (each lies inside a single Source tag)

They are the small charges a person really does incur twice in a day: a toll
passed twice, a fixed international-payment fee, two ride-hailing trips. A set
holds each triple once, so it drops every genuine repeat and eats the next
repeat toll on sight.

A real card export proved the same point on the incoming side — it carried an
exact twin (same day, merchant and amount) twice. Nothing but multiplicity
separates them. tests/test_model.py pins both with synthetic rows.

Balance WOULD tell every row apart and is still the wrong key: parse.py sets balance=None on every credit-card row because FNB
writes the literal string "0" there, and the CREDIT CARD sheet has no Balance
column at all. Keeping Balance out of the identity also leaves it free to serve
as an independent verification signal via verify.check_chain().

═══ WHY INTEGER CENTS, NOT str(amount) ═══
openpyxl returns int for whole rands and float for everything else, so on a
real sheet more than half of the amounts stringify to something the parser's
Decimal never emits:

    sheet  int(-3)      str -> '-3'        CSV  Decimal('-3.00')   -> '-3.00'
    sheet  float(-615.3) str -> '-615.3'   CSV  Decimal('-615.30') -> '-615.30'

Re-importing a real cheque export with a str() key matched fewer than half of
its rows and would have re-inserted the rest AS DUPLICATES. That failure is loud and
immediate; the set failure is silent and deferred. Both are fatal.

Float arithmetic drifts even though the stored values are clean (summing a whole
sheet of two-decimal amounts as float leaves a tail such as ...00000118), so
money is normalised through Decimal and compared as int cents.
"""

from __future__ import annotations

from collections import Counter
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import openpyxl

TXNS_SHEET = "TRANSACTIONS"
FIRST_DATA_ROW = 4

# Column order in TRANSACTIONS: Date | Description | Category | Amount | Balance | Source
COL_DATE = 0
COL_DESCRIPTION = 1
COL_AMOUNT = 3
COL_SOURCE = 5

COLUMNS = ("Date", "Description", "Category", "Amount", "Balance", "Source")

# D12: the card lives on its own sheet, in its own layout — a title block, then
# a header row, then Date | Description | Category | Amount. No Balance, no
# Source. Found by its header, not by a row number.
CARD_SHEET = "CREDIT CARD"
CARD_HEADER = ("Date", "Description", "Category", "Amount")

# The writer substitutes this when a description is empty (blank-narrative bank
# fees carry it). The parser emits "". Without this mapping every one of them
# re-imports.
NO_DESCRIPTION = "(no description)"

_CENTS = Decimal("0.01")


class ModelReadError(Exception):
    """Raised rather than skipping a row. A silently dropped row is a row that
    re-imports as a duplicate."""


def amount_cents(value) -> int:
    """Normalise money to integer cents, whatever the source type.

    Accepts int (sheet, whole rands), float (sheet), Decimal (parser) or str.
    Routing through Decimal(str(value)) is deliberate: Decimal(float) would
    carry the binary expansion into the key.
    """
    try:
        dec = value if isinstance(value, Decimal) else Decimal(str(value))
        if not dec.is_finite():
            raise ModelReadError(f"amount is not a finite number: {value!r}")
        # quantize() used to sit outside this guard, so Decimal("1E+26") escaped
        # as a raw decimal.InvalidOperation — bypassing the row-number wrapper
        # in load_existing and leaving the operator with no idea which row.
        return int(dec.quantize(_CENTS, rounding=ROUND_HALF_UP) * 100)
    except ModelReadError:
        raise
    except (InvalidOperation, ValueError, TypeError, ArithmeticError) as exc:
        raise ModelReadError(f"amount is not a usable number: {value!r}") from exc


def _compose(d, description, amount) -> tuple:
    """The dedupe identity. Three parts, and no more."""
    if isinstance(d, datetime):
        d = d.date()
    if not isinstance(d, date):
        raise ModelReadError(f"date is not a date: {d!r}")

    # NFC-normalise before comparing. The real sheet is entirely NFC today (its
    # only non-ASCII characters are the U+2014 em dashes in category labels), but any
    # macOS-originated or PDF-derived text arrives NFD, and "CAFÉ" composed
    # would not match "CAFÉ" decomposed. Normalising cannot merge two distinct
    # FNB strings; it can only stop one string failing to match itself.
    text = unicodedata.normalize("NFC", (description or "")).strip()
    if text == NO_DESCRIPTION:
        text = ""
    # Internal whitespace is NOT collapsed. FNB descriptions are fixed-width —
    # merchant [0:20], card mask [20:31], posting date [33:39] — and many real
    # rows carry runs of 2+ spaces that distinguish merchants. Collapsing
    # them would manufacture collisions that do not exist in the data.
    return (d.isoformat(), text, amount_cents(amount))


def key(txn) -> tuple:
    """The identity of a parsed transaction. Reads date, description, amount."""
    return _compose(txn.date, txn.description, txn.amount)


def row_hash(txn) -> str:
    """key(), hashed: how a private list names a real row without holding its
    date, description or amount in the clear (D13's hold list)."""
    import hashlib

    d, text, cents = key(txn)
    return hashlib.sha256(f"{d}|{text}|{cents}".encode("utf-8")).hexdigest()


def load_existing(xlsx_path: str) -> Counter:
    """Every row already in TRANSACTIONS, counted.

    read_only=True, data_only=True, and the workbook is closed before returning.
    Nothing here writes.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if TXNS_SHEET not in wb.sheetnames:
            raise ModelReadError(f"{xlsx_path} has no {TXNS_SHEET} sheet")
        ws = wb[TXNS_SHEET]
        # In read_only mode openpyxl trusts the sheet's <dimension> tag, which is
        # written by the producer and is routinely wrong outside Excel. A sheet
        # declaring A1:F6 over 10 rows silently returned 3 keys, after which
        # the other 7 rows re-imported as duplicates — silent and unbounded.
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        seen: Counter = Counter()
        for offset, row in enumerate(
            ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True)
        ):
            if len(row) <= COL_SOURCE:
                raise ModelReadError(
                    f"row {FIRST_DATA_ROW + offset} of {TXNS_SHEET} has "
                    f"{len(row)} columns, expected at least {COL_SOURCE + 1} "
                    f"({', '.join(COLUMNS)})"
                )
            if row[COL_DATE] is None and row[COL_AMOUNT] is None:
                # Excel omits empty rows from the XML, but a trailing blank can
                # still surface through iter_rows. A row with neither a date nor
                # an amount is not a transaction.
                continue
            try:
                seen[
                    _compose(row[COL_DATE], row[COL_DESCRIPTION], row[COL_AMOUNT])
                ] += 1
            except ModelReadError as exc:
                raise ModelReadError(
                    f"row {FIRST_DATA_ROW + offset} of {TXNS_SHEET}: {exc}"
                ) from exc
        return seen
    finally:
        wb.close()


def new_rows(parsed: list, existing: Counter) -> list:
    """The rows in `parsed` that the sheet does not already account for.

    Dedupes against the MODEL only — never within the batch. Two identical rows
    in one export are two real payments (spec §5.12: every such group on a real
    ledger reconciles against the balance chain).

    `existing` is not mutated, so run.py can report before it writes and the
    same Counter can be reused. Order is preserved: Format A is newest-first
    with non-monotonic dates, and the balance chain — not the date — defines
    the sequence.

    Does NOT filter on is_transaction. That is the caller's job, immediately
    before writing; keeping it out here means a CR.INT.RATE row still consumes
    its counterpart on the sheet instead of looking perpetually new.
    """
    remaining = existing.copy()
    fresh = []
    for txn in parsed:
        k = key(txn)
        if remaining[k] > 0:
            remaining[k] -= 1
        else:
            fresh.append(txn)
    return fresh


@dataclass
class CardRow:
    """One card transaction, for analysis only — card rows are never written.
    `source` says where it came from: the sheet, an export, or both."""

    date: date
    description: str
    category: str
    amount: Decimal
    source: str
    account: str = "card"
    is_transaction: bool = True
    held: bool = False
    row: int | None = None  # the sheet row, for rows read off the sheet

    @property
    def merchant(self) -> str:
        return self.description[:24].strip()


def _card_date(value) -> date:
    return value.date() if isinstance(value, datetime) else value


def load_card_sheet(xlsx_path: str) -> list:
    """The CREDIT CARD sheet's rows as CardRow, read-only. [] when the workbook
    has no such sheet. Rows without a date (totals, notes) are not purchases.

    D12, 2026-09-27: the pipeline analysed TRANSACTIONS only, so card spend was
    visible nowhere — the cheque side sees only the settlement, a transfer.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if CARD_SHEET not in wb.sheetnames:
            return []
        ws = wb[CARD_SHEET]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
        header = next(
            (
                i
                for i, r in enumerate(rows[:30])
                if tuple(str(c or "").strip() for c in r[:4]) == CARD_HEADER
            ),
            None,
        )
        if header is None:
            raise ModelReadError(f"{CARD_SHEET} has no Date / Description / Category / Amount header")
        out = []
        for offset, r in enumerate(rows[header + 1 :], header + 2):
            if len(r) < 4 or not isinstance(r[0], (date, datetime)):
                continue
            try:
                amount = Decimal(amount_cents(r[3])) / 100
            except ModelReadError as exc:
                raise ModelReadError(f"row {offset} of {CARD_SHEET}: {exc}") from exc
            out.append(
                CardRow(
                    date=_card_date(r[0]),
                    description=str(r[1] or "").strip(),
                    category=str(r[2] or "").strip(),
                    amount=amount,
                    source="sheet",
                    row=offset,
                )
            )
        return out
    finally:
        wb.close()


def truncated_card_cells(sheet: list, incoming: list) -> list:
    """(cell, sheet text, full text) for every CREDIT CARD description an older
    import cut at a comma: the same purchase (date and amount) is in an export,
    and the export's description carries on past the sheet's at a comma. The
    parser is right and the sheet is wrong, so writer.fix_text_cells() fixes
    the cells; the code is not bent to match them."""
    column = "ABCD"[CARD_HEADER.index("Description")]
    out, claimed = [], set()
    for s in sheet:
        for i, e in enumerate(incoming):
            full = (e.description or "").strip()
            if (
                i not in claimed
                and s.row is not None
                and _card_date(e.date) == _card_date(s.date)
                and amount_cents(e.amount) == amount_cents(s.amount)
                and s.description
                and full.upper().startswith(s.description.upper())
                and full[len(s.description):].startswith(",")  # cut AT a comma
            ):
                claimed.add(i)
                out.append((f"{column}{s.row}", s.description, full))
                break
    return out


def _prefix_compatible(a: str, b: str) -> bool:
    """An older import CUT some sheet descriptions at a comma (the parser keeps
    the whole line), so it is enough that one description starts the other."""
    a, b = (a or "").strip().upper(), (b or "").strip().upper()
    return bool(a and b) and (a.startswith(b) or b.startswith(a))


def merge_card_rows(sheet: list, incoming: list) -> list:
    """Sheet rows and incoming card-export rows as one list, each purchase once.

    A sheet row and an export row are the same purchase when the date and the
    amount agree and one description starts the other. Each export row is
    claimed at most once, so a genuine repeat (two coffees, same price, same
    day) survives. A category typed on the sheet wins over the rules; the
    export's description wins, because it is never truncated.
    """
    claimed, out = set(), []
    for s in sheet:
        match = None
        for i, e in enumerate(incoming):
            if i in claimed:
                continue
            if (
                _card_date(e.date) == _card_date(s.date)
                and amount_cents(e.amount) == amount_cents(s.amount)
                and _prefix_compatible(s.description, e.description)
            ):
                match = i
                break
        if match is None:
            out.append(CardRow(s.date, s.description, s.category, Decimal(s.amount), "sheet"))
            continue
        claimed.add(match)
        e = incoming[match]
        typed = s.category if s.category not in ("", "Uncategorised") else ""
        out.append(
            CardRow(
                _card_date(e.date),
                e.description.strip(),
                typed or getattr(e, "category", "") or s.category,
                Decimal(amount_cents(e.amount)) / 100,
                "both",
            )
        )
    for i, e in enumerate(incoming):
        if i not in claimed:
            out.append(
                CardRow(
                    _card_date(e.date),
                    e.description.strip(),
                    getattr(e, "category", ""),
                    Decimal(amount_cents(e.amount)) / 100,
                    "export",
                )
            )
    return out


def drop_key(txn) -> tuple:
    """Identity of a row INSIDE one drop: the dedupe triple plus the balance.

    Balance belongs here and not in key(): two exports of the same account print
    the same running balance for the same row, so it separates an overlap copy
    (identical balance) from a genuine repeat (a new chain position). The sheet
    comparison cannot use it — see the module docstring.
    """
    balance = None if txn.balance is None else amount_cents(txn.balance)
    return key(txn) + (balance,)


def dedupe_drop(files: list) -> tuple:
    """Merge the parsed files of ONE account into a single list, each row once.

    D1, 2026-09-24. Two exports with overlapping date ranges are the normal way
    people export, and dedupe used to run against the workbook only — so every
    overlapping row was "new" twice. Measured on a real overlapping pair: every
    row in the overlap would have been appended twice.

    Multiset UNION: each drop_key appears max(count in any one file) times, never
    the sum. A genuine repeat inside one file (two R27 tolls on one day, at two
    chain positions) survives, and an overlap copy collapses.

    Order: the first file's rows in file order, then each later file's surplus
    in its own order. Returns (merged, dropped_count).
    """
    merged: list = []
    have: Counter = Counter()
    dropped = 0
    for rows in files:
        seen_here: Counter = Counter()
        for txn in rows:
            k = drop_key(txn)
            seen_here[k] += 1
            if seen_here[k] > have[k]:
                have[k] += 1
                merged.append(txn)
            else:
                dropped += 1
    return merged, dropped
